import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import DateEntry
import csv, os
from datetime import datetime
import matplotlib.pyplot as plt
from collections import defaultdict
from reportlab.pdfgen import canvas
import platform
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

FILE = None
USER_FILE = "users.csv"
#==================TITLE ========================================================================================================
root = tk.Tk()
root.title("Zaki,akbar,iskandar: PERSONAL EXPENSE TRACKING SYSTEM")
root.state("zoomed")
root.configure(bg="#0b1220")

#==============ANIMATION LOGIN==================================================================================================
root.attributes("-alpha", 0)

def fade_in():
    alpha = root.attributes("-alpha")

    if alpha < 1:
        alpha += 0.02
        root.attributes("-alpha", alpha)
        root.after(20, fade_in)

fade_in()

logged_user = ""
backup_data = []

def get_user_file():
    return f"{logged_user}.csv"



# ================= USER =======================================================================================================
def save_user(u, p):
    with open(USER_FILE, "a", newline="") as f:
        csv.writer(f).writerow([u, p, "user"])

def user_exists(u):
    if not os.path.exists(USER_FILE):
        return False
    for r in csv.reader(open(USER_FILE)):
        if len(r) >= 2 and r[0] == u:
            return True
    return False

def check_login(u, p):
    if not os.path.exists(USER_FILE):
        return False, None

    for r in csv.reader(open(USER_FILE)):
        if len(r) >= 3 and r[0] == u and r[1] == p:
            return True, r[2]   # role
    return False, None

# ================= SWITCH =====================================================================================================
def show_register():
    login_card.pack_forget()
    register_card.pack(expand=True)

def back_login():
    register_card.pack_forget()
    login_card.pack(expand=True)

# ================= REGISTER ====================================================================================================
def register():
    u = reg_user.get()
    p = reg_pass.get()

    if len(u) == 0 or len(p) == 0:
        messagebox.showerror("Error", "Please fill all fields")
        return

    if len(u) > 15:
        messagebox.showerror("Error", "Username max 15 characters")
        return

    if len(p) < 6:
        messagebox.showerror("Error", "Password min 6 characters")
        return

    if user_exists(u):
        messagebox.showerror("Error", "Username already exists")
        return

    save_user(u, p)
    messagebox.showinfo("Success", "Account created!")
    back_login()

# ================= LOGIN ======================================================================================================
def login():

    global logged_user, logged_pass

    success, role = check_login(
        user.get().strip(),
        password.get().strip()
    )

    if success:

        logged_user = user.get().strip()

        if not os.path.exists(get_user_file()):
          open(get_user_file(), "w").close()

        logged_pass = password.get().strip()

        login_frame.destroy()
        dashboard()

    else:
        messagebox.showerror(
            "Login Failed",
            "Wrong username or password"
        )

# ================= LOGIN UI ===================================================================================================
login_frame = tk.Frame(root, bg="#0b1220")
login_frame.pack(expand=True, fill="both")

login_card = tk.Frame(login_frame, bg="#111b2e", padx=40, pady=35)
login_card.place(relx=0.5, rely=0.5, anchor="center")

tk.Label(login_card, text="LOGIN",
         fg="#4cc9f0", bg="#111b2e",
         font=("Arial", 18, "bold")).pack(pady=10)

tk.Label(login_card, text="Username", fg="white", bg="#111b2e").pack(anchor="w")
def check_username(new_value):
    if len(new_value) > 15:
        messagebox.showwarning("Limit", "Username max 15 characters only!")
        return False
    return True

vcmd_user = root.register(check_username)

user = tk.Entry(login_card, width=25, validate="key", validatecommand=(vcmd_user, "%P"))
user.pack(pady=5)

tk.Label(login_card, text="Password", fg="white", bg="#111b2e").pack(anchor="w")
pass_frame = tk.Frame(login_card, bg="#111b2e")
pass_frame.pack(pady=5)

# VALIDATION FUNCTION
def check_password(new_value):
    if new_value == "":
        return True

    if not new_value.isdigit():
        messagebox.showwarning("Invalid", "Password must be numbers only!")
        return False

    if len(new_value) > 6:
        messagebox.showwarning("Limit", "Password max 6 digits only!")
        return False

    return True

vcmd_pass = root.register(check_password)

# PASSWORD ENTRY (ONLY 1 ENTRY)
password = tk.Entry(
    pass_frame,
    width=20,
    show="*",
    validate="key",
    validatecommand=(vcmd_pass, "%P")
)
password.pack(side="left")

# HOLD FUNCTION
def show_password(event):
    password.config(show="")

def hide_password(event):
    password.config(show="*")

# BUTTON (TEPI)
login_eye_btn = tk.Button(pass_frame, text="🔒", width=2)
login_eye_btn.pack(side="left", padx=5)

login_eye_btn.bind("<ButtonPress-1>", show_password)
login_eye_btn.bind("<ButtonRelease-1>", hide_password)

tk.Button(login_card, text="LOGIN", bg="#22c55e",
          fg="white", width=25, command=login).pack(pady=10)

tk.Button(login_card, text="CREATE NEW ACCOUNT",
          bg="#3b82f6", fg="white",
          width=25, command=show_register).pack()

# ================= REGISTER UI (FIXED UX) =====================================================================================
register_card = tk.Frame(login_frame, bg="#111b2e", padx=40, pady=30)

tk.Label(register_card,
         text="CREATE NEW ACCOUNT",
         fg="#4cc9f0",
         bg="#111b2e",
         font=("Arial", 18, "bold")).pack(pady=10)

form = tk.Frame(register_card, bg="#111b2e")
form.pack(pady=10)

# USERNAME
tk.Label(form, text="Username (max 15 chars):",
         fg="white", bg="#111b2e").grid(row=0, column=0, sticky="w", pady=5)

def check_reg_username(new_value):
    if len(new_value) > 15:
        messagebox.showwarning("Limit", "Username max 15 characters only!")
        return False
    return True

vcmd_reg_user = root.register(check_reg_username)

reg_user = tk.Entry(
    form,
    width=18,
    validate="key",
    validatecommand=(vcmd_reg_user, "%P")
)

reg_user.grid(row=0, column=1, pady=5, padx=10)

# PASSWORD
tk.Label(form, text="Password (max 6 numbers):",
         fg="white", bg="#111b2e").grid(row=1, column=0, sticky="w", pady=5)

# FRAME (ENTRY SAHAJA, NO BUTTON)
reg_pass_frame = tk.Frame(form, bg="#111b2e")
reg_pass_frame.grid(row=1, column=1, pady=5, padx=10)
def check_reg_password(new_value):
    if new_value == "":
        return True

    if not new_value.isdigit():
        messagebox.showwarning("Invalid", "Password must be numbers only!")
        return False

    if len(new_value) > 6:
        messagebox.showwarning("Limit", "Password max 6 digits only!")
        return False

    return True

vcmd_reg_pass = root.register(check_reg_password)
# PASSWORD (VISIBLE)
reg_pass = tk.Entry(
    reg_pass_frame,
    width=18,
    validate="key",
    validatecommand=(vcmd_reg_pass, "%P")
)
reg_pass.pack(side="left")

# INFO BOX (CLEAR EXPLANATION)
info = tk.Label(register_card,
         text="✔ Username: max 15 characters (letters/numbers/symbols)\n"
              "✔ Password: max 6 numbers\n"
              "✔ Must remember for login",
         fg="#cbd5e1", bg="#111b2e", justify="left")
info.pack(pady=10)

btns = tk.Frame(register_card, bg="#111b2e")
btns.pack(pady=10)

tk.Button(btns, text="CREATE ACCOUNT",
          bg="#22c55e", fg="white",
          width=20, command=register).grid(row=0, column=0, padx=5)

tk.Button(btns, text="BACK",
          bg="#6b7280", fg="white",
          width=20, command=back_login).grid(row=0, column=1, padx=5)

register_card.pack_forget()

# ================= VALIDATION (NUMBERS ONLY) =====================================================================================

def only_number(char):
    return char.isdigit() or char == ""

validate_cmd = root.register(only_number)

# =========================================================== DASHBOARD =========================================================
def dashboard():

    global tree, total_label, backup_data

    backup_data = []  # FIX: initialize backup

    category = tk.StringVar(value="Select Category")
    rm = tk.StringVar()
    sen = tk.StringVar()
    desc = tk.StringVar()
    search = tk.StringVar()
    search_cat = tk.StringVar()
    salary_var = tk.StringVar()

    def parse_row(r):
        """
        Tukar CSV row ke format standard TreeView
        """
        return [
            r[1],  # date
            r[2],  # day
            r[3],  # category
            r[4],  # amount
            r[5],  # description
        ]

    def refresh_data():
        tree.delete(*tree.get_children())

        if not os.path.exists(get_user_file()):
            total()
            finance_status()
            return

        data = []

        with open(get_user_file(), newline="", encoding="utf-8") as f:
            for r in csv.reader(f):
                if len(r) < 6:
                    continue

                if r[0] != logged_user:
                    continue

                data.append(r)

        # SORT TARIKH
        try:
            data.sort(key=lambda x: datetime.strptime(x[1], "%Y-%m-%d"))
        except:
            pass

        # DISPLAY
        for r in data:
            tree.insert("", "end", values=r[1:])

        total()
        finance_status()

    def sort_by_date():

        items = []

        for i in tree.get_children():
            values = tree.item(i)["values"]

            if not values or values[0] == "TOTAL":
                continue

            try:
                date_obj = datetime.strptime(
                    str(values[0]),
                    "%Y-%m-%d"
                )

                items.append((date_obj, values))

            except:
                pass

        items.sort(key=lambda x: x[0])   # lama → baru

        tree.delete(*tree.get_children())

        for _, row in items:
            tree.insert("", "end", values=row)

   
    # ================= VALIDATION DIGIT =================
    def validate_input(char):
        return char.isdigit() or char == ""

    validate_cmd = root.register(validate_input)

    # ================= BACKUP =================

    def backup():
     snapshot = [
        tree.item(i)["values"]
        for i in tree.get_children()
     ]
     backup_data.append(snapshot)

    def undo():

        if not backup_data:
            messagebox.showinfo("Undo", "No more undo history!")
            return

        last_state = backup_data.pop()

        tree.delete(*tree.get_children())

        for row in last_state:
            tree.insert("", "end", values=row)

        save()
        excel(auto=True)
        total()
        finance_status()
    # ================= SAVE =================

    def save():

        sort_by_date()

        with open(get_user_file(), "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)

            for item in tree.get_children():
                values = tree.item(item)["values"]

                if values[0] == "TOTAL":
                    continue

                writer.writerow(values)

    # ================= LOAD =================

    def load():
        tree.delete(*tree.get_children())

        if not os.path.exists(get_user_file()):
            return

        data = []

        with open(get_user_file(), newline="", encoding="utf-8") as f:
            for row in csv.reader(f):

                if len(row) < 6:
                    continue

                if row[0].strip() != logged_user.strip():
                    continue

                data.append(row)

        # SORT IKUT TARIKH
        data.sort(
            key=lambda x: datetime.strptime(x[1], "%Y-%m-%d")
        )

        # MASUKKAN KE TABLE
        for row in data:
         tree.insert(
            "",
            "end",
            values=parse_row(row),
            tags=(row[3],)
        )

        total()
        finance_status()

    # ================= ADD =================
    def add():
        backup()

        if category.get() == "Select Category":
            messagebox.showerror("Error", "Please select category")
            return

        if not rm.get() and not sen.get():
            messagebox.showerror("Error", "Please enter amount")
            return

        try:
            amount = float(rm.get() or 0) + float(sen.get() or 0) / 100
        except:
            messagebox.showerror("Error", "Invalid amount")
            return

        if amount <= 0:
            messagebox.showerror("Error", "Amount must be greater than 0")
            return

        d = date.get_date()

        # SAVE CSV (ONLY USER DATA)
        with open(get_user_file(), "a", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([
                logged_user,
                d.strftime("%Y-%m-%d"),
                d.strftime("%A"),
                category.get(),
                f"{amount:.2f}",
                desc.get()
            ])

        load()
        excel(auto=True)
        total()
        finance_status()

        # reset form
        category.set("Select Category")
        rm.set("")
        sen.set("")
        desc.set("")
        date.set_date(datetime.now())

    # ================= DELETE =================
    def delete():
        backup()
        for i in tree.selection():
            tree.delete(i)
        save()
        excel(auto=True)
        total()
        finance_status()

    # ================= CLEAR =================
    def clear():
        confirm = messagebox.askyesno(
            "Confirm Clear",
            "Are you sure you want to clear all data?"
        )

        if not confirm:
            return

        backup()
        tree.delete(*tree.get_children())
        save()
        excel(auto=True)
        total()
        finance_status()


    def total():
        t = 0

        # buang TOTAL lama sahaja
        for item in tree.get_children():
            values = tree.item(item)["values"]
            if values and values[0] == "TOTAL":
                tree.delete(item)

        # kira total
        for item in tree.get_children():
            values = tree.item(item)["values"]

            if not values or values[0] == "TOTAL":
                continue

            try:
                t += float(str(values[3]).replace("RM", "").strip())
            except:
                pass

        total_label.config(text=f"Total Spending: RM {t:.2f}")

        # insert TOTAL bawah sekali
        tree.insert(
            "",
            "end",
            values=("TOTAL", "", "", "", f"RM {t:.2f}", ""),
            tags=("total_row",)
        )
    # ================= FINANCE STATUS =================
    def finance_status():

        try:
            salary = float(salary_var.get())
        except:
            salary = 0

        spent = 0

        for i in tree.get_children():
            values = tree.item(i)["values"]

            try:
                if values[0] == "TOTAL":
                    continue

                # AMOUNT = index 3
                raw = str(values[3]).replace("RM", "").replace(",", "").strip()

                if raw != "":
                    spent += float(raw)

            except:
                continue

        balance = salary - spent

        if salary == 0:
            status = "Enter salary first"
        elif balance < 0:
            status = "⚠ OVERSPENDING"
        elif balance < salary * 0.3:
            status = "⚠ HIGH SPENDING"
        else:
            status = "✅ HEALTHY"

        status_label.config(
            text=f"{status} | Balance: RM {balance:.2f}"
        )
    # ================= SEARCH =================
    def search_data():

        q = search.get().lower()
        cat = search_cat.get()
        dfilter = cal.get_date().strftime("%Y-%m-%d") if cal_used.get() else None

        tree.delete(*tree.get_children())

        data = []
        total_search = 0

        if not os.path.exists(get_user_file()):
            return

        with open(get_user_file(), newline="", encoding="utf-8") as f:

            for r in csv.reader(f):

                if len(r) < 6:
                    continue

                if r[0] == "TOTAL":
                    continue

                # FILTER SEARCH TEXT
                ok = True

                if q and q not in str(r).lower():
                    ok = False

                # CATEGORY FILTER
                if cat != "Select Category":
                    if cat.lower() != r[3].lower():
                        ok = False

                # DATE FILTER
                if dfilter and dfilter != r[1]:
                    ok = False

                if ok:
                    data.append(r)

                    try:
                        total_search += float(
                            str(r[4]).replace("RM", "").strip()
                        )
                    except:
                        pass

        # SORT BY DATE
        try:
            data.sort(
                key=lambda x: datetime.strptime(x[1], "%Y-%m-%d")
            )
        except:
            pass

        # DISPLAY DATA (remove username only)
        for r in data:
            tree.insert(
                "",
                "end",
                values=parse_row(r),
                tags=(r[3],)
            )

        # TOTAL ROW
        tree.insert(
            "",
            "end",
            values=("TOTAL", "", "", "", f"RM {total_search:.2f}", ""),
            tags=("total_row",)
        )

        total_label.config(
            text=f"Total Spending: RM {total_search:.2f}"
        )

        excel(auto=True)

    # ================= LOAD ALL =================
    def load_all():
        tree.delete(*tree.get_children())

        data = []

        if os.path.exists(get_user_file()):
            for r in csv.reader(open(get_user_file(), newline="", encoding="utf-8")):

                if len(r) < 6:
                    continue

                if r[0].strip() != logged_user.strip():
                    continue

                data.append(r)

        # SORT IKUT TARIKH
        try:
            data.sort(
                key=lambda x: datetime.strptime(x[1], "%Y-%m-%d")
            )
        except:
            pass

        for r in data:
         tree.insert(
            "",
            "end",
            values=parse_row(r),
            tags=(r[3],)
        )

        # reset filter
        search.set("")
        search_cat.set("Select Category")
        cal_used.set(0)
        cal.config(state="disabled")

        total()
        excel(auto=True)
        finance_status()
    # ================= CHART =================
    def chart():

        if len(tree.get_children()) == 0:
            messagebox.showwarning(
                "No Data",
                "Please add data first."
            )
            return

        data = defaultdict(float)

        for i in tree.get_children():

            r = tree.item(i)["values"]

            try:
                amount = float(
                    str(r[3])
                    .replace("RM", "")
                    .replace(",", "")
                    .strip()
                )

                data[str(r[2])] += amount

            except:
                continue

        if not data:
            messagebox.showwarning(
                "No Valid Data",
                "Cannot generate chart."
            )
            return

        plt.figure(figsize=(6,6))
        plt.title("Spending Breakdown")

        plt.pie(
            data.values(),
            labels=data.keys(),
            autopct="%1.1f%%"
        )

        plt.show()

        # ================= PDF (CLEAN REPORT STYLE) =================
    def pdf():

        file = f"{logged_user}_report.pdf"
        c = canvas.Canvas(file)

        # ================= TITLE =================
        c.setFillColorRGB(0.15, 0.35, 0.75)
        c.setFont("Helvetica-Bold", 18)
        c.drawCentredString(300, 800, "PERSONAL EXPENSE REPORT")

        # USER & DATE
        c.setFillColorRGB(0, 0, 0)
        c.setFont("Helvetica", 10)

        c.drawString(
            50,
            775,
            f"User : {logged_user}"
        )

        c.drawString(
            50,
            760,
            f"Generated : {datetime.now().strftime('%A, %d %B %Y, %I:%M %p')}"
        )

        c.line(40, 750, 560, 750)

        # ================= TABLE HEADER =================
        y = 720

        c.setFillColorRGB(0.20, 0.45, 0.85)
        c.rect(40, y, 520, 25, fill=1)

        c.setFillColorRGB(1, 1, 1)
        c.setFont("Helvetica-Bold", 10)

        c.drawString(50, y + 8, "Date")
        c.drawString(120, y + 8, "Day")
        c.drawString(200, y + 8, "Category")
        c.drawString(300, y + 8, "Amount")
        c.drawString(400, y + 8, "Description")

        y -= 25

        total_spending = 0
        row_num = 0

        # ================= DATA =================
        for item in tree.get_children():

            r = tree.item(item)["values"]

            if not r:
                continue

            if str(r[0]) == "TOTAL":
                continue

            # zebra color
            if row_num % 2 == 0:
                c.setFillColorRGB(0.95, 0.95, 0.95)
            else:
                c.setFillColorRGB(1, 1, 1)

            c.rect(40, y, 520, 20, fill=1)

            c.setFillColorRGB(0, 0, 0)
            c.setFont("Helvetica", 9)

            c.drawString(50, y + 5, str(r[0]))
            c.drawString(120, y + 5, str(r[1]))
            c.drawString(200, y + 5, str(r[2]))
            c.drawString(300, y + 5, str(r[3]))
            c.drawString(400, y + 5, str(r[4])[:20])

            try:
                total_spending += float(
                    str(r[3]).replace("RM", "").strip()
                )
            except:
                pass

            row_num += 1
            y -= 20

            # NEW PAGE
            if y < 70:

                c.showPage()

                y = 780

                c.setFillColorRGB(0.20, 0.45, 0.85)
                c.rect(40, y, 520, 25, fill=1)

                c.setFillColorRGB(1, 1, 1)
                c.setFont("Helvetica-Bold", 10)

                c.drawString(50, y + 8, "Date")
                c.drawString(120, y + 8, "Day")
                c.drawString(200, y + 8, "Category")
                c.drawString(300, y + 8, "Amount")
                c.drawString(400, y + 8, "Description")

                y -= 25

        # ================= TOTAL =================
        y -= 15

        c.setFillColorRGB(1, 0.90, 0.25)
        c.rect(40, y, 520, 25, fill=1)

        c.setFillColorRGB(0, 0, 0)
        c.setFont("Helvetica-Bold", 11)

        c.drawString(
            50,
            y + 8,
            f"TOTAL SPENDING : RM {total_spending:.2f}"
        )

        # ================= FOOTER =================
        c.setFont("Helvetica-Oblique", 8)

        c.drawCentredString(
            300,
            30,
            "Generated by Personal Expense Tracking System"
        )

        c.save()

        messagebox.showinfo(
            "PDF",
            f"PDF Report Saved As\n{file}"
        )

        if platform.system() == "Windows":
            os.startfile(file)
        # ================= excel (CLEAN REPORT STYLE) =================
    def excel(auto=True):

        file = f"{logged_user}_report.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Expense Report"

        # ================= TITLE =================

        ws.merge_cells("A1:E1")

        title = ws["A1"]
        title.value = "PERSONAL EXPENSE REPORT"

        title.font = Font(
            name="Arial",
            size=22,
            bold=True,
            color="FFFFFF"
        )

        title.fill = PatternFill(
            "solid",
            fgColor="1E3A8A"
        )

        title.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        ws.row_dimensions[1].height = 35

        # ================= USER INFO =================

        ws["A3"] = "User"
        ws["B3"] = logged_user

        ws["A4"] = "Generated"

        ws["B4"] = datetime.now().strftime(
            "%A, %d %B %Y, %I:%M %p"
        )

        ws["A3"].font = Font(bold=True)
        ws["A4"].font = Font(bold=True)

        # ================= HEADER =================

        headers = [
            "Date",
            "Day",
            "Category",
            "Amount",
            "Description"
        ]

        header_row = 6

        ws.row_dimensions[header_row].height = 25

        blue_fill = PatternFill(
            "solid",
            fgColor="2563EB"
        )

        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        for col, header in enumerate(headers, 1):

            cell = ws.cell(
                row=header_row,
                column=col
            )

            cell.value = header

            cell.fill = blue_fill

            cell.font = Font(
                color="FFFFFF",
                bold=True,
                size=11
            )

            cell.border = border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # ================= DATA =================


        row = 7
        total_spending = 0

        light_blue = PatternFill(
            "solid",
            fgColor="EFF6FF"
        )

        white_fill = PatternFill(
            "solid",
            fgColor="FFFFFF"
        )

        center_align = Alignment(
            horizontal="center",
            vertical="center"
        )

        left_align = Alignment(
            horizontal="left",
            vertical="center"
        )

        for item in tree.get_children():

            values = tree.item(item)["values"]

            if not values:
                continue

            if values[0] == "TOTAL":
                continue

            for col, value in enumerate(values, 1):

                cell = ws.cell(
                    row=row,
                    column=col
                )

                cell.value = value
                cell.border = border

                # warna selang-seli
                if row % 2 == 0:
                    cell.fill = light_blue
                else:
                    cell.fill = white_fill

                # alignment
                if col == 5:
                    cell.alignment = left_align
                else:
                    cell.alignment = center_align

            try:
                total_spending += float(
                    str(values[3])
                    .replace("RM", "")
                    .strip()
                )
            except:
                pass

            row += 1
       # ================= TOTAL =================

        total_row = row + 2

        gold_fill = PatternFill(
            "solid",
            fgColor="FFD700"
        )

        total_font = Font(
            bold=True,
            size=12,
            color="000000"
        )

        thick_border = Border(
            left=Side(style="medium"),
            right=Side(style="medium"),
            top=Side(style="medium"),
            bottom=Side(style="medium")
        )

        # merge A sampai C
        ws.merge_cells(
            start_row=total_row,
            start_column=1,
            end_row=total_row,
            end_column=3
        )

        title_cell = ws.cell(
            total_row,
            1
        )

        title_cell.value = "TOTAL SPENDING"
        title_cell.fill = gold_fill
        title_cell.font = total_font
        title_cell.alignment = Alignment(
            horizontal="center"
        )

        # amount
        amount_cell = ws.cell(
            total_row,
            4
        )

        amount_cell.value = f"RM {total_spending:.2f}"
        amount_cell.fill = gold_fill
        amount_cell.font = total_font
        amount_cell.alignment = Alignment(
            horizontal="center"
        )

        # description kosong
        empty_cell = ws.cell(
            total_row,
            5
        )

        empty_cell.fill = gold_fill

        # border semua cell total
        for col in range(1, 6):

            ws.cell(
                total_row,
                col
            ).border = thick_border

        # ================= COLUMN WIDTH =================

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 40

        # ================= FREEZE HEADER =================

        ws.freeze_panes = "A7"

        wb.save(file)

        if not auto:

            messagebox.showinfo(
                "Excel",
                f"Excel Report Saved\n{file}"
            )

            if platform.system() == "Windows":
                os.startfile(file)
        
    # ================= UI =================
    main = tk.Frame(root, bg="#0b1220")
    main.pack(fill="both", expand=True)

    # LEFT PANEL
    left = tk.Frame(main, bg="#111b2e", padx=20, pady=20)
    left.pack(side="left", fill="y")

    tk.Label(left, text=f"USER: {logged_user}",
             fg="#4cc9f0", bg="#111b2e",
             font=("Arial", 12, "bold")).pack()

    tk.Label(left, text="Monthly Salary", fg="white", bg="#111b2e").pack(anchor="w")
    salary_var = tk.StringVar()
    def validate_salary(value):

     if value == "":
        return True

     try:
        float(value)
        return True
     except:
        return False

    salary_validate = root.register(validate_salary)

    tk.Entry(
        left,
        textvariable=salary_var,
        validate="key",
        validatecommand=(salary_validate, "%P")
    ).pack(fill="x")

    tk.Button(left, text="CHECK STATUS",
              bg="#3b82f6", fg="white",
              command=finance_status).pack(fill="x", pady=5)

    status_label = tk.Label(left, text="Status: -",
                            fg="#facc15", bg="#111b2e",
                            wraplength=200)
    status_label.pack()

    # FORM
    form = tk.Frame(left, bg="#0f172a", padx=10, pady=10)
    form.pack(fill="x", pady=10)

    tk.Label(form, text="Category", fg="white", bg="#0f172a").pack(anchor="w")

    cat_box = ttk.Combobox(
        form,
        textvariable=category,
        values=["Select Category","Food","Transport","Bills","Shopping","Utilities","Entertainment","Health","Other"],
        state="readonly"
    )
    cat_box.pack(fill="x")
    cat_box.set("Select Category")

    tk.Label(form, text="Date", fg="white", bg="#0f172a").pack(anchor="w")

    date = DateEntry(form,date_pattern="yyyy-mm-dd",state="readonly")
    date.pack(fill="x")

    
    def block_typing(event):
        if event.keysym not in ("Down", "Up", "Return"):
            return "break"

    date.bind("<Key>", block_typing)

    tk.Label(form, text="Amount", fg="white", bg="#0f172a").pack(anchor="w")

    amt_frame = tk.Frame(form, bg="#0f172a")
    amt_frame.pack(fill="x")

    tk.Label(amt_frame, text="RM", fg="white", bg="#0f172a").pack(side="left")
    tk.Entry(amt_frame, textvariable=rm, width=10,
             validate="key", validatecommand=(validate_cmd, "%S")).pack(side="left")

    tk.Label(amt_frame, text="SEN", fg="white", bg="#0f172a").pack(side="left")
    tk.Entry(amt_frame, textvariable=sen, width=5,
             validate="key", validatecommand=(validate_cmd, "%S")).pack(side="left", padx=5)

    tk.Label(form, text="Description", fg="white", bg="#0f172a").pack(anchor="w")
    tk.Entry(form, textvariable=desc).pack(fill="x")

    tk.Button(left, text="➕ ADD", bg="#22c55e", command=add).pack(fill="x", pady=5)
    tk.Button(left, text="↩ UNDO", bg="#f59e0b", command=undo).pack(fill="x", pady=5)
    tk.Button(left, text="🗑 CLEAR", bg="#ef4444", command=clear).pack(fill="x", pady=5)

    # RIGHT PANEL
    right = tk.Frame(main, bg="#0b1220")
    right.pack(side="right", fill="both", expand=True)

    topbar = tk.Frame(right, bg="#111b2e", padx=10, pady=10)
    topbar.pack(fill="x")


    tk.Entry(topbar, textvariable=search, width=20).pack(side="left")

    search_cat = tk.StringVar()
    search_cat.set("Select Category")

    search_cat_box = ttk.Combobox(
        topbar,
        textvariable=search_cat,
        values=["Select Category","Food","Transport","Bills","Shopping","Utilities","Entertainment","Health","Other"],
        state="readonly"   
    )

    search_cat_box.pack(side="left", padx=5)
    search_cat_box.set("Select Category")


        # ================= DATE FILTER =================

    cal_used = tk.IntVar(value=0)

    def toggle_date():
        if cal_used.get():
            cal.config(state="readonly")
        else:
            cal.config(state="disabled")

    chk = ttk.Checkbutton(
        topbar,
        text="Date Filter",
        variable=cal_used,
        command=toggle_date
    )

    chk.pack(side="left", padx=5)

    cal = DateEntry(
        topbar,
        state="disabled",
        date_pattern="yyyy-mm-dd"
    )

    cal.pack(side="left", padx=5)


    tk.Button(topbar, text="SEARCH", bg="#3b82f6",
              fg="white", command=search_data).pack(side="left")
    
    tk.Button(topbar, text="VIEW ALL", bg="#f59e0b",
              fg="white", command=load_all).pack(side="left", padx=5)
    
    tk.Button(topbar, text="CHART", bg="#8b5cf6",
              fg="white", command=chart).pack(side="right", padx=5)

    tk.Button(topbar, text="PDF", bg="#10b981",
              fg="white", command=pdf).pack(side="right", padx=5)
    '''tk.Button(topbar,text="EXCEL",bg="#16a34a",fg="white",command=excel).pack(side="right", padx=5)'''
    tk.Button(topbar,text="EXCEL", bg="#16a34a",fg="white",command=lambda: excel(auto=False)).pack(side="right", padx=5)

    tk.Button(topbar, text="DELETE", bg="#ef4444",
              fg="white", command=delete).pack(side="right", padx=5)

    # TABLE
    tree = ttk.Treeview(right,
                        columns=("Date","Day","Category","Amount","Description"),
                        show="headings")
    style = ttk.Style()

    style.theme_use("winnative")  

    style.configure(
        "Treeview.Heading",
        background="#3700FF",   
        foreground="white",    
        font=("Arial", 10, "bold")
    )
    tree.pack(fill="both", expand=True)

    tree.tag_configure(
    "total_row",
    background="#facc15",
    font=("Arial", 10, "bold")
    )

    style = ttk.Style()

    style.configure(
        "Treeview",
        rowheight=30,
        borderwidth=1,
        relief="solid"
    )

    style.configure(
        "Treeview.Heading",
        relief="raised",
        borderwidth=1
    )

    tree.tag_configure("Food", background="#edf5ec")
    tree.tag_configure("Transport", background="#dbeafe")
    tree.tag_configure("Bills", background="#fee2e2")
    tree.tag_configure("Shopping", background="#fef3c7")
    tree.tag_configure("Entertainment", background="#ede9fe")
    tree.tag_configure("Health", background="#fce7f3")
    tree.tag_configure("Utilities", background="#e0f2fe")
    tree.tag_configure("Other", background="#f3f4f6")

    '''tree.tag_configure("Food", background="#90EE90")
    tree.tag_configure("Transport", background="#87CEEB")
    tree.tag_configure("Bills", background="#FFB6B6")
    tree.tag_configure("Shopping", background="#FFE599")
    tree.tag_configure("Entertainment", background="#D8BFD8")
    tree.tag_configure("Health", background="#FF5AE6")
    tree.tag_configure("Utilities", background="#B0E0E6")
    tree.tag_configure("Other", background="#D3D3D3")'''


    # COLUMN SETTINGS
    for c in tree["columns"]:
        tree.heading(c, text=c)
        tree.column(c, anchor="center", width=130)

    total_label = tk.Label(
        right,
        text="Total: RM 0.00",
        fg="#facc15",
        bg="#0b1220",
        font=("Arial", 12, "bold")
    )
    total_label.pack()


    load()
    excel(auto=True)
    total()
    finance_status()

root.mainloop()